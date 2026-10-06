"""Review endpoints, run against the real graph with the offline model backend.

Each run is a real review, so its ids and timings are its own. Responses are therefore checked for
the *shape* of the committed contract examples -- a renamed or dropped field still fails -- and for
behaviour directly: what a submission stores, who may decide a run, what the export contains.
"""

from __future__ import annotations

from io import BytesIO

import pytest
from openpyxl import load_workbook

from mlr import service
from tests.api.conftest import SAMPLE, assert_same_shape, fixture

SUBMISSION = {"tenant_id": "acme_pharma", "product_id": "zelvoran"}
DECISION = {"decision": "reject", "reviewer": "a.reviewer", "tier": "senior_regulatory"}


# --- Service-level -------------------------------------------------------------------------------


def test_every_documented_endpoint_is_routed(client):
    """The contract table in docs/contracts.md, as a test: nothing listed there may be missing."""
    paths = client.get("/openapi.json").json()["paths"]
    documented = [
        ("post", "/reviews"),
        ("get", "/reviews"),
        ("get", "/reviews/{run_id}"),
        ("get", "/queue"),
        ("post", "/reviews/{run_id}/decision"),
        ("get", "/reviews/{run_id}/audit"),
        ("get", "/reviews/{run_id}/report.json"),
        ("get", "/reviews/{run_id}/report.xlsx"),
        ("get", "/tenants"),
        ("get", "/tenants/{tenant_id}"),
        ("get", "/tenants/{tenant_id}/rules"),
        ("get", "/tenants/{tenant_id}/rules/{rule_id}"),
        ("post", "/tenants/{tenant_id}/rules"),
        ("put", "/tenants/{tenant_id}/rules/{rule_id}"),
        ("post", "/tenants/{tenant_id}/rules/{rule_id}/test"),
        ("post", "/tenants/{tenant_id}/rules/{rule_id}/activate"),
        ("post", "/tenants/{tenant_id}/rules/{rule_id}/retire"),
        ("post", "/tenants/{tenant_id}/rules/export"),
        ("post", "/evals/run"),
        ("get", "/evals/runs"),
        ("get", "/evals/runs/{eval_run_id}"),
        ("post", "/evals/sweep"),
        ("get", "/graph"),
    ]
    missing = [f"{m.upper()} {p}" for m, p in documented if m not in paths.get(p, {})]
    assert not missing, f"documented endpoints with no route: {missing}"


def test_health_says_which_components_are_stubbed(client):
    """A demo serving example data has to be able to say so on screen."""
    body = client.get("/health").json()
    assert body["status"] == "ok"
    assert body["stubs"]["review_service"] is service.IS_STUB
    assert body["stubs"]["eval_runner"] is True
    assert "not medical, legal or regulatory advice" in body["disclaimer"]


def test_graph_matches_the_contract_example(client):
    assert client.get("/graph").json() == fixture("GET_graph.json")


def test_no_endpoint_lets_the_system_approve_content(client):
    """The only way a run completes is a human decision. There is no approve route to call."""
    paths = client.get("/openapi.json").json()["paths"]
    assert not [p for p in paths if "approve" in p.lower()]


# --- Submitting ----------------------------------------------------------------------------------


def test_submitting_text_registers_a_run_that_can_be_read_at_once(client, real_runs):
    response = client.post("/reviews", data={**SUBMISSION, "text": "Zelvoran works."})
    assert response.status_code == 202
    assert_same_shape(response.json(), fixture("POST_reviews.json"))
    assert response.json()["status"] == "queued"
    assert client.get(f"/reviews/{response.json()['run_id']}").status_code == 200


def test_the_stored_document_is_the_text_the_checks_read(client, submitted):
    """Front matter gone and emphasis markers stripped, so finding offsets point at the right words."""
    document = client.get(f"/reviews/{submitted}/document").json()
    assert document["filename"] == SAMPLE.name
    assert document["document"].startswith("# Nearly 6 in 10 adults reached PASI-90")
    assert "synthetic: true" not in document["document"]
    assert "**" not in document["document"]
    # Audience and channel are recorded as submitted, never inferred from the front matter.
    assert (document["audience"], document["channel"]) == ("hcp", "sales_aid")

    report = client.get(f"/reviews/{submitted}").json()
    quoted = [f for f in report["findings"] if f["quote"]]
    assert quoted, "the sample should produce findings that quote the piece"
    for finding in quoted:
        assert finding["quote"] in document["document"]
        start, end = finding["char_span"]
        if end:  # not every producer records an offset; when one is given it must be right
            assert document["document"][start:end] == finding["quote"]


def test_the_submission_reaches_the_service_with_its_audience_and_channel(client, monkeypatch):
    """They decide which tenant rules are in scope, so dropping them would change the review."""
    seen = {}

    def fake_start_run(**kwargs):
        seen.update(kwargs)
        return "run-xyz"

    monkeypatch.setattr(service, "start_run", fake_start_run)
    response = client.post(
        "/reviews",
        data={**SUBMISSION, "text": "Copy.", "audience": "patient", "channel": "leaflet"},
    )
    assert response.json()["run_id"] == "run-xyz"
    assert (seen["audience"], seen["channel"], seen["document"]) == ("patient", "leaflet", "Copy.")


@pytest.mark.parametrize(
    ("data", "files", "status"),
    [
        ({**SUBMISSION}, None, 422),  # neither a file nor text
        ({**SUBMISSION, "text": "   "}, None, 422),  # blank text
        ({**SUBMISSION, "text": "x"}, {"file": ("a.md", b"y", "text/markdown")}, 422),  # both
        ({"tenant_id": "nobody", "product_id": "zelvoran", "text": "x"}, None, 404),
        ({"tenant_id": "acme_pharma", "product_id": "nexapril", "text": "x"}, None, 422),
        ({**SUBMISSION, "text": "x", "audience": "everyone"}, None, 422),
        ({**SUBMISSION}, {"file": ("a.docx", b"y", "application/octet-stream")}, 415),
        ({**SUBMISSION}, {"file": ("a.txt", b"\xff\xfe\x00bad", "text/plain")}, 422),
        ({**SUBMISSION}, {"file": ("a.md", b"---\ntitle: t\n---\n", "text/markdown")}, 422),
    ],
)
def test_a_bad_submission_is_refused_before_a_run_starts(client, monkeypatch, data, files, status):
    def must_not_start(**_kwargs):
        raise AssertionError("a refused submission must not start a run")

    monkeypatch.setattr(service, "start_run", must_not_start)
    assert client.post("/reviews", data=data, files=files).status_code == status


def test_a_pdf_upload_is_read_as_text(client, real_runs):
    import pymupdf

    with pymupdf.open() as pdf:
        pdf.new_page().insert_text((72, 72), "Zelvoran achieved PASI-75 in 68.4% of adults.")
        data = pdf.tobytes()

    response = client.post("/reviews", data=SUBMISSION, files={"file": ("ad.pdf", data)})
    assert response.status_code == 202
    stored = client.get(f"/reviews/{response.json()['run_id']}/document").json()
    assert "PASI-75 in 68.4%" in stored["document"]


# --- Reading -------------------------------------------------------------------------------------


def test_a_finished_run_has_the_contract_shape_and_waits_for_a_person(client, submitted):
    report = client.get(f"/reviews/{submitted}").json()
    assert_same_shape(report, fixture("GET_reviews_{run_id}.json"))
    assert report["status"] == "awaiting_human"
    assert report["route"], "a run that has been routed names who it went to"
    assert any(f["severity_final"] == "blocker" for f in report["findings"])


def test_an_unknown_run_is_a_404(client):
    for path in ("", "/audit", "/report.json", "/report.xlsx", "/document"):
        assert client.get(f"/reviews/nope{path}").status_code == 404
    assert client.post("/reviews/nope/decision", json=DECISION).status_code == 404


def test_a_run_id_that_could_escape_the_data_directory_is_never_looked_up(client, monkeypatch):
    def must_not_be_called(run_id):
        raise AssertionError(f"looked up an unsafe run id: {run_id!r}")

    monkeypatch.setattr(service, "get_status", must_not_be_called)
    assert client.get("/reviews/..%5C..%5Csecrets").status_code == 404
    assert client.get("/reviews/a.b").status_code == 404


def test_the_run_list_has_the_contract_shape_and_filters(client, submitted):
    listed = client.get("/reviews").json()
    assert_same_shape(listed, fixture("GET_reviews.json"))
    assert [r["run_id"] for r in listed["runs"]] == [submitted]
    assert client.get("/reviews", params={"tenant_id": "acme_pharma"}).json()["total"] == 1
    assert client.get("/reviews", params={"tenant_id": "borealis_bio"}).json() == {
        "runs": [],
        "total": 0,
    }
    assert client.get("/reviews", params={"status": "complete"}).json()["total"] == 0
    assert client.get("/reviews", params={"status": "approved"}).status_code == 422


def test_the_queue_has_the_contract_shape_and_filters_by_tier(client, submitted):
    tier = client.get(f"/reviews/{submitted}").json()["route"]
    queue = client.get("/queue", params={"tier": tier}).json()
    assert_same_shape(queue, fixture("GET_queue.json"))
    assert [t["run_id"] for t in queue["tasks"]] == [submitted]
    assert client.get("/queue", params={"tier": "nobody"}).json() == {"tier": "nobody", "tasks": []}
    assert len(client.get("/queue").json()["tasks"]) == 1


def test_the_audit_trail_has_the_contract_shape_and_is_complete(client, submitted):
    trail = client.get(f"/reviews/{submitted}/audit").json()
    assert_same_shape(trail, fixture("GET_reviews_{run_id}_audit.json"))
    assert trail["integrity"]["sequence_complete"] is True
    nodes = [e["node"] for e in trail["events"]]
    assert nodes[0] == "ingest" and "grounding_gate" in nodes and "route" in nodes


def test_audit_integrity_reports_a_dropped_or_duplicated_event(client, submitted, monkeypatch):
    """The check has to be able to fail, or it is decoration."""
    events = service.get_audit(submitted)
    broken = [e for e in events if e["seq"] != 3] + [events[0]]
    monkeypatch.setattr(service, "get_audit", lambda _run_id: broken)

    integrity = client.get(f"/reviews/{submitted}/audit").json()["integrity"]
    assert integrity["sequence_complete"] is False
    assert integrity["duplicate_seqs"] == [1]


# --- Deciding ------------------------------------------------------------------------------------


def _decision_for(client, run_id: str, **overrides) -> dict:
    """A valid decision from the tier the run was actually routed to."""
    return {**DECISION, "tier": client.get(f"/reviews/{run_id}").json()["route"], **overrides}


def test_a_decision_resumes_the_run_to_completion(client, submitted):
    response = client.post(f"/reviews/{submitted}/decision", json=_decision_for(client, submitted))
    assert response.status_code == 200, response.text
    assert_same_shape(response.json(), fixture("POST_reviews_{run_id}_decision.json"))
    assert response.json()["status"] == "complete"
    assert response.json()["route"], "even a completed run names who was responsible"
    assert client.get("/queue").json()["tasks"] == []
    assert client.get(f"/reviews/{submitted}/audit").json()["events"][-1]["node"] == "finalize"


@pytest.mark.parametrize(
    "payload",
    [
        {"decision": "approve"},  # unattributed
        {**DECISION, "decision": "auto_approve"},  # not a decision a human can make
        {**DECISION, "per_finding": [{"finding_id": "f1", "verdict": "maybe"}]},
        {**DECISION, "approved_by_system": True},  # unknown field
    ],
)
def test_an_invalid_decision_never_reaches_the_run(client, monkeypatch, payload):
    def must_not_resume(*_args):
        raise AssertionError("an invalid decision must not resume a run")

    monkeypatch.setattr(service, "resume", must_not_resume)
    assert client.post("/reviews/anyrun/decision", json=payload).status_code == 422


def test_a_reviewer_of_the_wrong_tier_cannot_decide_an_escalated_run(client, submitted):
    """Otherwise a standard reviewer could clear a piece that policy escalated past them."""
    routed_to = client.get(f"/reviews/{submitted}").json()["route"]
    assert routed_to != "standard_reviewer", "the sample is expected to escalate"

    response = client.post(
        f"/reviews/{submitted}/decision",
        json=_decision_for(client, submitted, tier="standard_reviewer"),
    )
    assert response.status_code == 403
    assert routed_to in response.json()["detail"]
    assert client.get(f"/reviews/{submitted}").json()["status"] == "awaiting_human"


def test_a_run_that_has_been_decided_cannot_be_decided_again(client, submitted):
    decision = _decision_for(client, submitted)
    assert client.post(f"/reviews/{submitted}/decision", json=decision).status_code == 200
    assert client.post(f"/reviews/{submitted}/decision", json=decision).status_code == 409


# --- Exporting -----------------------------------------------------------------------------------


def test_the_json_export_has_one_located_row_per_finding(client, submitted):
    report = client.get(f"/reviews/{submitted}").json()
    exported = client.get(f"/reviews/{submitted}/report.json").json()
    assert_same_shape(exported, fixture("GET_reviews_{run_id}_report.json"))
    assert [r["finding_id"] for r in exported["rows"]] == [f["id"] for f in report["findings"]]
    # A finding that quotes the piece is placed under the heading it actually sits under.
    located = {r["location"] for r in exported["rows"] if r["quote"]}
    assert located and "" not in located


def test_the_export_leaves_location_blank_rather_than_guessing(client, submitted):
    """With no document to look in, or nothing quoted, there is no location to report."""
    from api.export import export_rows
    from mlr.schemas import ReviewReport

    report = ReviewReport.model_validate(client.get(f"/reviews/{submitted}").json())
    assert all(row["location"] == "" for row in export_rows(report, document=""))


def test_the_spreadsheet_has_the_same_rows_as_the_json_export(client, submitted):
    response = client.get(f"/reviews/{submitted}/report.xlsx")
    assert response.status_code == 200
    assert "spreadsheetml" in response.headers["content-type"]
    assert f"mlr-review-{submitted}.xlsx" in response.headers["content-disposition"]

    sheet = load_workbook(BytesIO(response.content))["Findings"]
    table = list(sheet.iter_rows(values_only=True))
    rows = client.get(f"/reviews/{submitted}/report.json").json()["rows"]

    assert len(table) == len(rows) + 1
    assert table[0][:4] == ("Finding", "Location", "Quote", "Issue")
    for exported, row in zip(table[1:], rows, strict=True):
        # An empty string is written as an empty cell, which reads back as None.
        assert ["" if v is None else v for v in exported] == list(row.values())


def test_exported_text_is_never_stored_as_a_spreadsheet_formula():
    """The quote column is untrusted upload text, and the file is meant to be passed around."""
    from api.export import export_xlsx
    from mlr.schemas import Finding, ReviewReport

    report = ReviewReport(
        run_id="r1",
        tenant_id="acme_pharma",
        product_id="zelvoran",
        status="awaiting_human",
        findings=[Finding(id="f1", check_type="business_rule", issue="x", quote="=1+1")],
    )
    sheet = load_workbook(BytesIO(export_xlsx(report)))["Findings"]
    quote = sheet.cell(row=2, column=3)
    assert quote.value == "=1+1"
    assert quote.data_type == "s", "a quote starting with '=' must be text, not a formula"
