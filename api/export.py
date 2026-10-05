"""The findings export: one row per finding, as JSON or as a spreadsheet.

Deterministic code. No model involved.

Both formats are built from the same rows, so the spreadsheet a reviewer circulates cannot disagree
with the JSON a downstream system ingests.
"""

from __future__ import annotations

from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from api.ingest import locate
from mlr.schemas import ReviewReport

DISCLAIMER = (
    "Synthetic data and fictional products. A demonstration only: not medical, legal or regulatory "
    "advice, and not for review of real promotional material."
)

# Column order of the export, and the header each one gets in the spreadsheet.
COLUMNS: list[tuple[str, str, int]] = [
    ("finding_id", "Finding", 10),
    ("location", "Location", 22),
    ("quote", "Quote", 48),
    ("issue", "Issue", 60),
    ("check", "Check", 20),
    ("severity", "Severity", 12),
    ("confidence", "Confidence", 12),
    ("grounded_by", "Grounded by", 14),
    ("citation_source", "Citation source", 34),
    ("citation_passage", "Citation passage", 60),
    ("suggested_fix", "Suggested fix", 60),
    ("rule", "Rule", 18),
    ("severity_rule", "Severity rule", 26),
]

_SEVERITY_FILL = {
    "blocker": PatternFill("solid", fgColor="F8D7DA"),
    "advisory": PatternFill("solid", fgColor="FFF3CD"),
}


def export_rows(report: ReviewReport, document: str = "") -> list[dict[str, Any]]:
    """One row per finding. Only the first citation is flattened into the row; all of them are in
    the report itself."""
    rows = []
    for f in report.findings:
        first = f.citations[0] if f.citations else None
        rows.append(
            {
                "finding_id": f.id,
                "location": locate(document, f.quote, f.char_span),
                "quote": f.quote,
                "issue": f.issue,
                "check": f.check_type,
                "severity": f.severity,
                "confidence": f.confidence,
                "grounded_by": f.grounded_by,
                "citation_source": first.source_id if first else "",
                "citation_passage": first.passage if first else "",
                "suggested_fix": f.suggested_fix,
                "rule": f"{f.rule_id} v{f.rule_version}" if f.rule_id else "",
                "severity_rule": f.severity_rule_id,
            }
        )
    return rows


def export_payload(report: ReviewReport, document: str = "") -> dict[str, Any]:
    return {
        "run_id": report.run_id,
        "tenant_id": report.tenant_id,
        "product_id": report.product_id,
        "policy_version": report.policy_version,
        "route": report.route,
        "rows": export_rows(report, document),
    }


def _write_text(cell: Any, value: Any) -> None:
    """Write a value, forcing strings to be stored as text.

    The quote column is copied verbatim from an uploaded document. A spreadsheet treats a cell that
    starts with ``=`` as a formula, so untrusted text written naively is a formula-injection hole in
    a file designed to be emailed around.
    """
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"


def export_xlsx(report: ReviewReport, document: str = "") -> bytes:
    """The same rows as :func:`export_payload`, as an .xlsx workbook."""
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Findings"

    for col, (_, header, width) in enumerate(COLUMNS, start=1):
        cell = sheet.cell(row=1, column=col, value=header)
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="top")
        sheet.column_dimensions[get_column_letter(col)].width = width

    for row_index, row in enumerate(export_rows(report, document), start=2):
        for col, (key, _, _) in enumerate(COLUMNS, start=1):
            cell = sheet.cell(row=row_index, column=col)
            _write_text(cell, row[key])
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if key == "severity" and row[key] in _SEVERITY_FILL:
                cell.fill = _SEVERITY_FILL[row[key]]
    sheet.freeze_panes = "A2"

    # Provenance travels with the file: a spreadsheet detached from the run that produced it should
    # still say which policy was in force and that a human, not the system, makes the decision.
    info = workbook.create_sheet("Run")
    for row_index, (label, value) in enumerate(
        [
            ("Run", report.run_id),
            ("Tenant", report.tenant_id),
            ("Product", report.product_id),
            ("Label version", report.label_version),
            ("Policy version", report.policy_version),
            ("Status", report.status),
            ("Routed to", report.route),
            ("Summary", report.summary_line()),
            ("Note", "The system recommends; a human decides."),
            ("Disclaimer", DISCLAIMER),
        ],
        start=1,
    ):
        info.cell(row=row_index, column=1, value=label).font = Font(bold=True)
        _write_text(info.cell(row=row_index, column=2), value)
    info.column_dimensions["A"].width = 18
    info.column_dimensions["B"].width = 100

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
